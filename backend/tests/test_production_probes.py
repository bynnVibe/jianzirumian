"""原生产评估的39项探针正式化；通过 run_isolated.py 执行，原始证据保持不变。"""
import asyncio
import multiprocessing
import random
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from fastapi import HTTPException
from langchain_core.embeddings import Embeddings

from app import config
from app.core import db, metrics_export, observability
from app.services.gateway import keys, channels, proxy, settings as gateway_settings


@pytest.fixture(autouse=True)
def isolated_db(tmp_path, monkeypatch):
    assert "production-assessment" in str(config.DATA_DIR), "必须用隔离入口执行"
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "probe.db")
    monkeypatch.setattr(db, "_conn", None)
    connection = db.get_conn()
    yield
    if db._conn is not None and db._conn is not connection:
        db._conn.close()
    connection.close()


def call_proxy(method, path, **kwargs):
    from app.llm_proxy import app
    async def request():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://probe") as client:
            return await client.request(method, path, **kwargs)
    return asyncio.run(request())


def add_channel(name="demo", priority=0, models=None):
    return channels.create_channel({"name": name, "base_url": f"https://{name}.invalid/v1",
                                    "models": models or ["test-model"], "priority": priority})


def body(stream=False):
    return {"model": "test-model", "messages": [{"role": "user", "content": "合成测试问题"}], "stream": stream}


@pytest.mark.parametrize("path,expected", [("/health", 200), ("/metrics", 200), ("/v1/models", 401)])
def test_proxy_http_health_metrics_and_auth(path, expected):
    assert call_proxy("GET", path).status_code == expected


@pytest.mark.parametrize("payload,status", [({}, 400), ({"model": "test-model"}, 401)])
def test_proxy_http_input_rejection(payload, status):
    assert call_proxy("POST", "/v1/chat/completions", json=payload).status_code == status


def test_proxy_models_filtered_by_key():
    add_channel(models=["allowed", "denied"])
    key = keys.create_key("test", allowed_models=["allowed"])
    response = call_proxy("GET", "/v1/models", headers={"Authorization": f"Bearer {key['key']}"})
    assert response.status_code == 200
    assert [m["id"] for m in response.json()["data"]] == ["allowed"]


@pytest.mark.parametrize("state,code", [("disabled", "disabled"), ("expired", "expired"), ("quota", "quota_exceeded")])
def test_key_rejection_states(state, code):
    key = keys.create_key("state", request_quota=1)
    if state == "disabled":
        keys.update_key(key["id"], {"enabled": False})
    elif state == "expired":
        keys.update_key(key["id"], {"expires_at": "2000-01-01T00:00:00"})
    else:
        keys.consume_request(key["id"])
    with pytest.raises(keys.KeyRejected) as exc:
        keys.authenticate(key["key"])
    assert exc.value.code == code


def test_key_plaintext_not_returned_after_creation():
    key = keys.create_key("secret")
    assert "key" not in keys.get_key(key["id"])
    assert db.query_one("SELECT key_hash FROM gateway_keys WHERE id=?", (key["id"],))["key_hash"] != key["key"]


def test_gateway_audit_block_without_upstream():
    gateway_settings.save_settings({"audit_mode": "block", "blacklist": ["PROBE_BLOCK_MARKER"]})
    key = keys.create_key("blocked")
    payload = body()
    payload["messages"][0]["content"] = "PROBE_BLOCK_MARKER"
    with pytest.raises(proxy.GatewayError) as exc:
        proxy.prepare(payload, key["key"])
    assert exc.value.status_code == 403
    assert db.query_one("SELECT status FROM gateway_logs")["status"] == "blocked"


def test_gateway_priority_and_health():
    high = add_channel("high", 10)
    add_channel("low", 0)
    assert channels.select_candidates("test-model")[0]["id"] == high["id"]
    for _ in range(3):
        channels.mark_failure(high["id"], "合成失败")
    assert channels.get_channel(high["id"])["health"] == "down"
    channels.mark_success(high["id"])
    assert channels.get_channel(high["id"])["health"] == "healthy"
    assert len(channels.select_candidates("test-model")) == 2


def test_gateway_json_failover_and_usage(monkeypatch):
    add_channel("bad", 10)
    add_channel("good", 0)
    key = keys.create_key("failover")
    ctx, seen, client_class = proxy.prepare(body(), key["key"]), [], httpx.AsyncClient
    def handler(request):
        seen.append(request.url.host)
        if request.url.host == "bad.invalid":
            return httpx.Response(503, json={"error": "合成上游失败"})
        return httpx.Response(200, json={"choices": [{"message": {"content": "合成回答"}}],
                                         "usage": {"prompt_tokens": 5, "completion_tokens": 3}})
    monkeypatch.setattr(proxy.httpx, "AsyncClient", lambda **kw: client_class(transport=httpx.MockTransport(handler), **kw))
    status, result = asyncio.run(proxy.json_response(ctx))
    assert status == 200 and result["choices"][0]["message"]["content"] == "合成回答"
    assert seen == ["bad.invalid", "good.invalid"]
    assert keys.get_key(key["id"])["token_used"] == 8
    assert keys.get_key(key["id"])["request_count"] == 1
    assert db.query_one("SELECT status,retries FROM gateway_logs") == {"status": "ok", "retries": 1}


@pytest.mark.parametrize("table,initial", [("wiki_compile_queue", "processing"), ("eval_runs", "running"), ("pipeline_runs", "running")])
def test_proxy_metrics_must_not_reset_agent_tasks(table, initial):
    now = datetime.now().isoformat()
    if table == "wiki_compile_queue":
        db.execute("INSERT INTO wiki_compile_queue(id,status,created_at,updated_at) VALUES ('probe',?,?,?)", (initial, now, now))
    elif table == "eval_runs":
        db.execute("INSERT INTO eval_runs(id,dataset_id,status,started_at) VALUES ('probe','synthetic',?,?)", (initial, now))
    else:
        db.execute("INSERT INTO pipeline_runs(id,status,started_at) VALUES ('probe',?,?)", (initial, now))
    old_connection, db._conn = db._conn, None
    assert call_proxy("GET", "/metrics").status_code == 200
    assert old_connection.execute(f"SELECT status FROM {table} WHERE id='probe'").fetchone()["status"] == initial


def quota_worker(database, raw_key, barrier, pipe):
    if db._conn:
        db._conn.close()
    db._conn, db.DB_PATH = None, Path(database)
    row = keys.authenticate(raw_key)
    barrier.wait(timeout=10)
    try:
        keys.consume_request(row["id"])
        pipe.send("accepted")
    except keys.KeyRejected as exc:
        pipe.send(exc.code)
    finally:
        db._conn.close()
        pipe.close()


def test_gateway_quota_atomic_across_processes():
    key = keys.create_key("concurrency", request_quota=1)
    context = multiprocessing.get_context("fork")
    barrier, pairs = context.Barrier(2), [context.Pipe() for _ in range(2)]
    workers = [context.Process(target=quota_worker, args=(str(db.DB_PATH), key["key"], barrier, child)) for _, child in pairs]
    try:
        for worker in workers:
            worker.start()
        results = []
        for parent, _ in pairs:
            assert parent.poll(20)
            results.append(parent.recv())
        for worker in workers:
            worker.join(timeout=15)
        assert all(worker.exitcode == 0 for worker in workers)
        assert sorted(results) == ["accepted", "quota_exceeded"]
    finally:
        for worker in workers:
            if worker.is_alive():
                worker.terminate()
            worker.join(5)
        for parent, child in pairs:
            parent.close()
            child.close()
    assert keys.get_key(key["id"])["request_count"] == 1


def test_proxy_metrics_must_not_duplicate_agent_trace():
    observability.start_trace("synthetic-agent", input_text="合成数据")
    observability.end_trace()
    assert 'jzrm_agent_traces_total{service="agent"} 1' in metrics_export.render_prometheus("agent")
    assert "jzrm_agent_traces_total" not in metrics_export.render_prometheus("llm-proxy")


def test_gateway_higher_weight_receives_more_traffic(monkeypatch):
    monkeypatch.setattr(channels.random, "random", random.Random(20261026).random)
    rows = [{"id": "small", "weight": 1}, {"id": "large", "weight": 9}]
    wins = sum(channels._weighted_order(rows)[0]["id"] == "large" for _ in range(2000))
    assert wins > 1400


@pytest.mark.parametrize("query,expected", [("看看这张图片", "ocr"), ("总结这份资料", "summarize"), ("保存到知识库", "upload")])
def test_tool_routing_positive(query, expected):
    from app.services.tool_orchestrator import _decide_action
    assert _decide_action(query) == expected


@pytest.mark.parametrize("query", ["不要上传，只识别文字", "不要保存到知识库，帮我总结", "只看看，不要入库"])
def test_tool_negative_intent_must_not_write(query):
    from app.services.tool_orchestrator import _decide_action
    assert _decide_action(query) != "upload"


def test_guest_attachment_cannot_select_write_tool(monkeypatch):
    from app.services import tool_orchestrator as tools
    from app.services.tools.base import ToolResult
    run = AsyncMock(return_value=ToolResult(tool_name="image_ocr", success=True, summary="合成识别"))
    write = AsyncMock(side_effect=AssertionError("不允许调用写工具"))
    monkeypatch.setattr(tools.image_ocr_tool, "run", run)
    monkeypatch.setattr(tools.image_upload_tool, "run", write)
    _, records = asyncio.run(tools.handle_attachments("保存到知识库", [{"path": "/synthetic.png", "filename": "synthetic.png"}], "guest", "", "s", "m"))
    assert records[0]["tool_name"] == "image_ocr"
    run.assert_awaited_once()
    write.assert_not_awaited()


class TinyEmbeddings(Embeddings):
    def embed_documents(self, texts):
        return [[1.0, 0.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, text):
        return [1.0, 0.0, 0.0, 0.0]


@pytest.fixture
def tiny_store(tmp_path, monkeypatch):
    from app.core import vector_store as vector_module
    monkeypatch.setattr(vector_module.EmbeddingFactory, "create", lambda *a, **kw: TinyEmbeddings())
    monkeypatch.setattr(config.settings, "VECTOR_STORE_PATH", str(tmp_path / "vectors/faiss_index"))
    store = vector_module.VectorStore()
    store.add_texts(["public apple", "alice apple", "bob apple"], [
        {"visibility": "public", "owner_id": "", "kb_id": "public"},
        {"visibility": "private", "owner_id": "alice", "kb_id": "alice"},
        {"visibility": "private", "owner_id": "bob", "kb_id": "bob"}])
    return store


def test_faiss_real_index_tenant_filter_and_reload(tiny_store):
    from app.core.vector_store import VectorStore
    for store in [tiny_store, VectorStore()]:
        found = store.similarity_search("apple", k=10, current_user_id="alice")
        assert {text for text, _, _ in found} == {"public apple", "alice apple"}
        public = store.similarity_search("apple", k=10, current_user_id="", kb_ids={"public"})
        assert [text for text, _, _ in public] == ["public apple"]


def test_bm25_real_index_tenant_filter(tiny_store, monkeypatch):
    from app.core import sparse_retriever as sparse_module
    monkeypatch.setattr(sparse_module, "vector_store", tiny_store)
    found = sparse_module.SparseRetriever().search("apple", 10, current_user_id="bob")
    assert {text for text, _, _ in found} == {"public apple", "bob apple"}


def test_docx_real_parser_preserves_table(tmp_path):
    from docx import Document
    from app.services.knowledge import KnowledgeService
    document = Document()
    document.add_paragraph("Synthetic paragraph")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Name", "Value"), ("Apple", "42")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / "synthetic.docx"
    document.save(path)
    text = KnowledgeService._extract_docx_text(str(path))
    assert "Synthetic paragraph" in text and "| Apple | 42 |" in text


def test_pdf_real_parser_preserves_page_numbers(tmp_path):
    import pymupdf
    from app.services.knowledge import KnowledgeService
    path = tmp_path / "synthetic.pdf"
    with pymupdf.open() as pdf:
        for text in ["Synthetic first page", "Synthetic second page"]:
            pdf.new_page().insert_text((72, 72), text)
        pdf.save(path)
    pages = KnowledgeService._extract_pdf_pages(str(path))
    assert [p.page_number for p in pages] == [1, 2]
    assert "first page" in pages[0].text and "second page" in pages[1].text


def test_bookmark_list_and_delete_ownership():
    from app.services import bookmark
    item = bookmark.create_bookmark("alice", "合成收藏", ["test"], "内容")
    assert len(bookmark.list_bookmarks("alice")) == 1
    assert bookmark.list_bookmarks("bob") == []
    assert bookmark.delete_bookmark("bob", item["id"]) is False
    assert bookmark.delete_bookmark("alice", item["id"]) is True


def test_kb_list_and_update_ownership():
    from app.services import kb
    kb.ensure_default_kb()
    private = kb.create_kb("测试私人库", "其他", "private", "alice", False)
    assert private["id"] in {item["id"] for item in kb.list_accessible_kbs("alice")}
    assert private["id"] not in {item["id"] for item in kb.list_accessible_kbs("bob")}
    with pytest.raises(PermissionError):
        kb.update_kb(private["id"], "越权修改", "其他", "bob", False)


@pytest.mark.parametrize("cache_hit", [False, True])
def test_private_file_authorization_before_cache(tmp_path, monkeypatch, cache_hit):
    from app.core.file_cache import file_cache, serve_cached_file
    from app.services.records import upload_records
    path = tmp_path / "private.txt"
    path.write_text("合成私有文件", encoding="utf-8")
    monkeypatch.setattr(upload_records, "get_records_by_source", lambda p: [{"visibility": "private", "owner_id": "alice"}])
    monkeypatch.setattr(file_cache, "get", lambda key: (b"cached", "text/plain") if cache_hit else None)
    with pytest.raises(HTTPException) as exc:
        asyncio.run(serve_cached_file(path, {"id": "bob", "role": "user"}))
    assert exc.value.status_code == 403
    response = asyncio.run(serve_cached_file(path, {"id": "alice", "role": "user"}))
    assert response.status_code == 200
    assert response.headers["X-File-Cache"] == ("HIT" if cache_hit else "MISS")


def test_redis_read_failure_degrades_and_cools_down(monkeypatch):
    from app.core.file_cache import FileCache
    cache = FileCache()
    def unavailable(*a, **kw):
        raise ConnectionError("合成 Redis 断线")
    monkeypatch.setattr(cache, "_get_client", lambda: SimpleNamespace(hget=unavailable))
    assert cache.get("synthetic-key") is None
    assert cache._failed_until > 0


def test_citation_range_validation():
    from app.core.evidence_checker import verify_citations
    assert verify_citations("[来源 1] [来源 0] [来源 4] [网络来源 99]", 2) == [0, 4]


def test_self_rag_provider_failure_is_nonfatal(monkeypatch):
    from app.core import evidence_checker
    monkeypatch.setattr(evidence_checker.LLMFactory, "create", lambda *a: (_ for _ in ()).throw(RuntimeError("合成模型失败")))
    assert asyncio.run(evidence_checker.check_evidence_sufficiency("问题", "合成上下文")) == (True, "")


@pytest.mark.parametrize("check", ["resource_closed", "terminal_status"])
def test_gateway_stream_cancel_contract(monkeypatch, check):
    add_channel()
    ctx = proxy.prepare(body(stream=True), keys.create_key("cancel")["key"])
    state, client_class = {"response_closed": False}, httpx.AsyncClient
    class Stream(httpx.AsyncByteStream):
        async def __aiter__(self):
            yield b'data: {"choices":[{"delta":{"content":"partial"}}]}\n\n'
            yield b'data: [DONE]\n\n'
        async def aclose(self):
            state["response_closed"] = True
    transport = httpx.MockTransport(lambda request: httpx.Response(200, stream=Stream()))
    monkeypatch.setattr(proxy.httpx, "AsyncClient", lambda **kw: client_class(transport=transport, **kw))
    async def cancel():
        generator = proxy.stream_response(ctx)
        assert "partial" in await anext(generator)
        await generator.aclose()
    asyncio.run(cancel())
    if check == "resource_closed":
        assert state["response_closed"] is True
    else:
        assert db.query_one("SELECT status FROM gateway_logs")["status"] == "cancelled"
